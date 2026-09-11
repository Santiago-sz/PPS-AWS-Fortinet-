"""
Tests para extractor.py.

Construye PDFs/DOCX válidos en memoria (sin depender de archivos externos)
para probar extracción real con pypdf/python-docx — nunca se mockean estas
librerías, se ejercita el parsing real sobre bytes válidos e inválidos.
"""

import io

import pytest
from docx import Document
from pypdf import PdfWriter

from extractor import (
    DOCX_CONTENT_TYPE,
    PDF_CONTENT_TYPE,
    ExtractionError,
    extract,
)


def _make_pdf_bytes(pages_text: list[str]) -> bytes:
    """Arma un PDF válido a mano (sin reportlab): un content stream con un
    solo `Tj` por página. Suficiente para que pypdf extraiga texto real."""
    n_pages = len(pages_text)
    objects = {}
    font_obj_num = 3
    page_obj_nums = list(range(4, 4 + n_pages))
    content_obj_nums = list(range(4 + n_pages, 4 + 2 * n_pages))

    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{n} 0 R" for n in page_obj_nums)
    objects[2] = f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode()
    objects[font_obj_num] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    for i, text in enumerate(pages_text):
        page_num = page_obj_nums[i]
        content_num = content_obj_nums[i]
        objects[page_num] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_obj_num} 0 R >> >> "
            f"/Contents {content_num} 0 R >>"
        ).encode()

        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream_content = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode()
        objects[content_num] = (
            f"<< /Length {len(stream_content)} >>\nstream\n".encode()
            + stream_content
            + b"\nendstream"
        )

    buf = bytearray()
    buf += b"%PDF-1.4\n"
    offsets = {}
    all_obj_nums = sorted(objects.keys())
    for num in all_obj_nums:
        offsets[num] = len(buf)
        buf += f"{num} 0 obj\n".encode()
        buf += objects[num]
        buf += b"\nendobj\n"

    xref_offset = len(buf)
    max_obj = max(all_obj_nums)
    buf += f"xref\n0 {max_obj + 1}\n".encode()
    buf += b"0000000000 65535 f \n"
    for num in range(1, max_obj + 1):
        if num in offsets:
            buf += f"{offsets[num]:010d} 00000 n \n".encode()
        else:
            buf += b"0000000000 00000 f \n"
    buf += b"trailer\n"
    buf += f"<< /Size {max_obj + 1} /Root 1 0 R >>\n".encode()
    buf += b"startxref\n"
    buf += f"{xref_offset}\n".encode()
    buf += b"%%EOF"
    return bytes(buf)


def _make_blank_pdf_bytes() -> bytes:
    """PDF válido de una página sin ningún content stream de texto —
    simula un PDF escaneado (imagen) sin capa de texto extraíble."""
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    bio = io.BytesIO()
    writer.write(bio)
    return bio.getvalue()


def _make_encrypted_pdf_bytes(password: str = "secret123") -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt(user_password=password)
    bio = io.BytesIO()
    writer.write(bio)
    return bio.getvalue()


def _make_docx_bytes(paragraphs: list[tuple[str, int | None]]) -> bytes:
    """paragraphs: lista de (texto, heading_level|None). heading_level=None
    -> párrafo normal; heading_level=N -> estilo 'Heading N'."""
    doc = Document()
    for text, level in paragraphs:
        if level is None:
            doc.add_paragraph(text)
        else:
            doc.add_heading(text, level=level)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _make_empty_docx_bytes() -> bytes:
    doc = Document()
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


class TestExtractPdf:
    def test_valid_pdf_extracts_full_text_and_page_metadata(self):
        pdf_bytes = _make_pdf_bytes(
            ["Articulo 1. Purpose of the policy.", "Articulo 2. Scope of application."]
        )

        text, metadata = extract(pdf_bytes, PDF_CONTENT_TYPE)

        assert "Articulo 1. Purpose of the policy." in text
        assert "Articulo 2. Scope of application." in text
        # Metadata trae un elemento por página, en orden, con char_span válido
        assert [m["page"] for m in metadata] == [1, 2]
        for m in metadata:
            assert 0 <= m["char_start"] < m["char_end"] <= len(text)
            assert text[m["char_start"] : m["char_end"]]

    def test_second_page_text_is_reachable_via_its_char_span(self):
        pdf_bytes = _make_pdf_bytes(["First page content here.", "Second page content here."])

        text, metadata = extract(pdf_bytes, PDF_CONTENT_TYPE)

        second = metadata[1]
        assert "Second page content here." in text[second["char_start"] : second["char_end"]]

    def test_corrupt_pdf_raises_extraction_error(self):
        with pytest.raises(ExtractionError) as exc_info:
            extract(b"this is not a pdf file at all, just garbage bytes", PDF_CONTENT_TYPE)

        assert exc_info.value.reason == "corrupt"

    def test_encrypted_pdf_raises_extraction_error(self):
        encrypted_bytes = _make_encrypted_pdf_bytes()

        with pytest.raises(ExtractionError) as exc_info:
            extract(encrypted_bytes, PDF_CONTENT_TYPE)

        assert exc_info.value.reason == "encrypted"

    def test_scanned_or_empty_pdf_raises_extraction_error(self):
        blank_bytes = _make_blank_pdf_bytes()

        with pytest.raises(ExtractionError) as exc_info:
            extract(blank_bytes, PDF_CONTENT_TYPE)

        assert exc_info.value.reason == "no_text_layer"

    def test_large_pdf_text_is_never_truncated(self):
        # MAX_DATA_CHARS en analyzer.py es 80_000 — probamos que extractor
        # nunca aplica un límite similar.
        big_paragraph = "Clausula de cumplimiento normativo de prueba. " * 3000
        pdf_bytes = _make_pdf_bytes([big_paragraph, big_paragraph])

        text, metadata = extract(pdf_bytes, PDF_CONTENT_TYPE)

        assert len(text) > 80_000
        assert "truncad" not in text.lower()
        assert text.count("Clausula de cumplimiento normativo de prueba.") == 6000


class TestExtractDocx:
    def test_valid_docx_extracts_text_and_heading_metadata(self):
        docx_bytes = _make_docx_bytes(
            [
                ("Politica de Seguridad", 1),
                ("Texto introductorio de la politica.", None),
                ("Articulo 1. Objetivo", 2),
                ("Contenido del articulo 1.", None),
            ]
        )

        text, metadata = extract(docx_bytes, DOCX_CONTENT_TYPE)

        assert "Politica de Seguridad" in text
        assert "Contenido del articulo 1." in text

        headings = [m for m in metadata if m["heading"] is not None]
        assert [h["heading"] for h in headings] == ["Politica de Seguridad", "Articulo 1. Objetivo"]
        assert [h["level"] for h in headings] == [1, 2]

    def test_docx_char_spans_align_with_full_text(self):
        docx_bytes = _make_docx_bytes(
            [("Articulo 1. Objetivo", 2), ("Contenido del articulo 1.", None)]
        )

        text, metadata = extract(docx_bytes, DOCX_CONTENT_TYPE)

        heading_meta = next(m for m in metadata if m["heading"] == "Articulo 1. Objetivo")
        assert text[heading_meta["char_start"] : heading_meta["char_end"]] == "Articulo 1. Objetivo"

    def test_corrupt_docx_raises_extraction_error(self):
        with pytest.raises(ExtractionError) as exc_info:
            extract(b"not a docx at all, just garbage", DOCX_CONTENT_TYPE)

        assert exc_info.value.reason == "corrupt"

    def test_empty_docx_raises_extraction_error(self):
        empty_bytes = _make_empty_docx_bytes()

        with pytest.raises(ExtractionError) as exc_info:
            extract(empty_bytes, DOCX_CONTENT_TYPE)

        assert exc_info.value.reason == "no_text_layer"


class TestExtractUnsupportedContentType:
    def test_unsupported_content_type_raises_extraction_error(self):
        with pytest.raises(ExtractionError) as exc_info:
            extract(b"whatever", "text/plain")

        assert exc_info.value.reason == "unsupported_content_type"
